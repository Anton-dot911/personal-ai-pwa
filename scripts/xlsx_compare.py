import sys, math
import openpyxl

# Використання: python3 scripts/xlsx_compare.py ЗРАЗОК.xlsx ЗГЕНЕРОВАНИЙ.xlsx
SRC = sys.argv[1]
GEN = sys.argv[2]

a = openpyxl.load_workbook(SRC)
b = openpyxl.load_workbook(GEN)
av = openpyxl.load_workbook(SRC, data_only=True)
bv = openpyxl.load_workbook(GEN, data_only=True)

results = []
def check(name, ok, detail=""):
    results.append((name, ok))
    print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  — {detail}" if detail and not ok else ""))

def argb(c):
    rgb = getattr(c, "rgb", None) if c is not None else None
    return rgb if isinstance(rgb, str) else None

def fill(cell):
    f = cell.fill
    return argb(f.fgColor) if f is not None and f.fill_type == "solid" else None

def same_value(x, y):
    if isinstance(x, float) or isinstance(y, float):
        try:
            return math.isclose(float(x), float(y), rel_tol=1e-9, abs_tol=1e-9)
        except (TypeError, ValueError):
            return False
    return x == y

check("аркуші і порядок", a.sheetnames == b.sheetnames, f"{a.sheetnames} vs {b.sheetnames}")

for name in a.sheetnames:
    wa, wb_ = a[name], b[name]
    check(f"[{name}] розмір {wa.max_row}x{wa.max_column}",
          (wa.max_row, wa.max_column) == (wb_.max_row, wb_.max_column),
          f"{(wb_.max_row, wb_.max_column)}")

    vals_bad = []
    fmt_bad = []
    style_bad = []
    cached_bad = []
    for row in wa.iter_rows(min_row=1, max_row=wa.max_row, max_col=wa.max_column):
        for ca in row:
            cb = wb_[ca.coordinate]
            if not same_value(ca.value, cb.value):
                vals_bad.append((ca.coordinate, ca.value, cb.value))
            if ca.number_format != cb.number_format:
                fmt_bad.append((ca.coordinate, ca.number_format, cb.number_format))
            sa = (fill(ca), argb(ca.font.color) if ca.font and ca.font.color else None,
                  bool(ca.font.b), float(ca.font.sz or 0), ca.font.name, bool(ca.alignment.wrap_text))
            sb = (fill(cb), argb(cb.font.color) if cb.font and cb.font.color else None,
                  bool(cb.font.b), float(cb.font.sz or 0), cb.font.name, bool(cb.alignment.wrap_text))
            if sa != sb:
                style_bad.append((ca.coordinate, sa, sb))
            va, vb = av[name][ca.coordinate].value, bv[name][ca.coordinate].value
            if isinstance(ca.value, str) and ca.value.startswith("=") and not same_value(va, vb):
                cached_bad.append((ca.coordinate, va, vb))
    check(f"[{name}] значення і формули", not vals_bad, f"{len(vals_bad)} відмінностей, напр. {vals_bad[:3]}")
    check(f"[{name}] формати чисел і дат", not fmt_bad, f"{len(fmt_bad)}, напр. {fmt_bad[:3]}")
    check(f"[{name}] стилі (заливка, шрифт, жирність, перенос)", not style_bad, f"{len(style_bad)}, напр. {style_bad[:2]}")
    check(f"[{name}] кешовані результати формул", not cached_bad, f"{len(cached_bad)}, напр. {cached_bad[:3]}")

    fa = wa.freeze_panes
    fb = wb_.freeze_panes
    check(f"[{name}] закріплення {fa}", fa == fb, f"{fb}")
    aa = wa.auto_filter.ref
    ab = wb_.auto_filter.ref
    check(f"[{name}] фільтр {aa}", aa == ab, f"{ab}")

    wid_bad = []
    for L, dim in wa.column_dimensions.items():
        if dim.width:
            other = wb_.column_dimensions[L].width
            if other is None or abs(other - dim.width) > 0.6:
                wid_bad.append((L, dim.width, other))
    check(f"[{name}] ширини стовпців", not wid_bad, f"{wid_bad[:5]}")

    def dv_key(dv):
        return (dv.type, dv.operator or 'between', str(dv.formula1), str(dv.formula2) if dv.formula2 is not None else None,
                str(dv.sqref), bool(dv.allow_blank), dv.error, dv.errorTitle)
    da = sorted(dv_key(d) for d in wa.data_validations.dataValidation)
    db = sorted(dv_key(d) for d in wb_.data_validations.dataValidation)
    check(f"[{name}] валідація даних ({len(da)} правил)", da == db,
          f"зразок {len(da)}, результат {len(db)}; відмінності: {[x for x in da if x not in db][:2]}")

    def cf_key(cf):
        return sorted((str(cf.sqref), r.type, tuple(r.formula or []),
                       argb(r.dxf.fill.bgColor) if r.dxf and r.dxf.fill else None,
                       argb(r.dxf.font.color) if r.dxf and r.dxf.font and r.dxf.font.color else None)
                      for r in cf.rules)
    ca_ = sorted(sum((cf_key(cf) for cf in wa.conditional_formatting), []))
    cb_ = sorted(sum((cf_key(cf) for cf in wb_.conditional_formatting), []))
    check(f"[{name}] умовне форматування ({len(ca_)} правил)", ca_ == cb_,
          f"відмінності: {[x for x in ca_ if x not in cb_][:2]}")

failed = [n for n, ok in results if not ok]
print()
print(f"ПІДСУМОК: {len(results) - len(failed)} з {len(results)} перевірок пройдено")
if failed:
    print("Не пройдено:", failed)
    sys.exit(1)
