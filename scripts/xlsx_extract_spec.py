import json, collections, datetime, re
import openpyxl
from openpyxl.utils import get_column_letter

# Використання: python3 scripts/xlsx_extract_spec.py ЗРАЗОК.xlsx ВИХІД.json
import sys
SRC = sys.argv[1]
OUT = sys.argv[2]
INPUT_FILLS = {"FFFFF6D6"}   # колір полів вводу в зразку

wb = openpyxl.load_workbook(SRC)
wbv = openpyxl.load_workbook(SRC, data_only=True)

def argb(color):
    rgb = getattr(color, "rgb", None) if color is not None else None
    return rgb if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000" else None

def fill_of(cell):
    f = cell.fill
    return argb(f.fgColor) if f is not None and f.fill_type == "solid" else None

def style_of(cell, dname, dsize):
    st = {}
    fl = fill_of(cell)
    if fl: st["fill"] = fl
    fc = argb(cell.font.color) if cell.font and cell.font.color else None
    if fc and fc != "FF000000": st["fontColor"] = fc
    if cell.font and cell.font.b: st["bold"] = True
    if cell.font and cell.font.name and cell.font.name != dname: st["fontName"] = cell.font.name
    if cell.font and cell.font.sz and float(cell.font.sz) != dsize: st["size"] = float(cell.font.sz)
    if cell.number_format and cell.number_format != "General": st["numFmt"] = cell.number_format
    if cell.alignment and cell.alignment.wrap_text: st["wrap"] = True
    return st

def enc(cell, cached):
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        obj = {"formula": v[1:]}
        if cached.value not in (None, ""): obj["result"] = cached.value
        return obj
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v

def num_or_str(x):
    if isinstance(x, (int, float)): return x
    s = str(x)
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else s

def sheet_spec(ws, wsv):
    max_r, max_c = ws.max_row, ws.max_column
    sizes, names = collections.Counter(), collections.Counter()
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
        for c in row:
            if c.font and c.font.sz: sizes[float(c.font.sz)] += 1
            if c.font and c.font.name: names[c.font.name] += 1
    default_size = sizes.most_common(1)[0][0] if sizes else 10.0
    font_name = names.most_common(1)[0][0] if names else "Calibri"

    columns = []
    for ci in range(1, max_c + 1):
        L = get_column_letter(ci)
        dim = ws.column_dimensions.get(L)
        width = float(dim.width) if dim is not None and dim.width else None
        h = style_of(ws.cell(1, ci), font_name, default_size)
        d = style_of(ws.cell(2, ci), font_name, default_size) if max_r >= 2 else {}
        col = {}
        if width: col["width"] = round(width, 2)
        if h: col["headerStyle"] = h
        if d: col["cellStyle"] = d
        if d.get("fill") in INPUT_FILLS: col["role"] = "input"
        columns.append(col)

    rows = [[enc(ws.cell(r, ci), wsv.cell(r, ci)) for ci in range(1, max_c + 1)] for r in range(1, max_r + 1)]

    overrides = []
    for r in range(1, max_r + 1):
        for ci in range(1, max_c + 1):
            cell = ws.cell(r, ci)
            st = style_of(cell, font_name, default_size)
            base = columns[ci - 1].get("headerStyle", {}) if r == 1 else columns[ci - 1].get("cellStyle", {})
            if st != base:
                overrides.append({"ref": cell.coordinate, "style": st})

    vals = []
    for dv in ws.data_validations.dataValidation:
        item = {"ref": str(dv.sqref), "type": dv.type,
                "formulae": [num_or_str(x) for x in (dv.formula1, dv.formula2) if x is not None]}
        if dv.operator and dv.operator != "between": item["operator"] = dv.operator
        if dv.allow_blank is not None: item["allowBlank"] = bool(dv.allow_blank)
        if dv.error: item["error"] = dv.error
        if dv.errorTitle: item["errorTitle"] = dv.errorTitle
        vals.append(item)

    conds = collections.OrderedDict()
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            rr = {"formula": rule.formula[0] if rule.formula else ""}
            if rule.dxf and rule.dxf.fill:
                fc = argb(rule.dxf.fill.bgColor) or argb(rule.dxf.fill.fgColor)
                if fc: rr["fill"] = fc
            if rule.dxf and rule.dxf.font and rule.dxf.font.color:
                fc = argb(rule.dxf.font.color)
                if fc: rr["fontColor"] = fc
            conds.setdefault(str(cf.sqref), []).append(rr)

    spec = {"name": ws.title, "font": {"name": font_name, "size": default_size},
            "columns": columns, "rows": rows}
    if overrides: spec["cellStyles"] = overrides
    if ws.freeze_panes: spec["freeze"] = ws.freeze_panes
    if ws.auto_filter.ref: spec["autoFilter"] = ws.auto_filter.ref
    if vals: spec["validations"] = vals
    if conds: spec["conditionals"] = [{"ref": k, "rules": v} for k, v in conds.items()]
    return spec, len(overrides)

spec = {"filename": "suppliers_template.xlsx", "sheets": []}
for ws in wb.worksheets:
    s, n_over = sheet_spec(ws, wbv[ws.title])
    spec["sheets"].append(s)
    n_formulas = sum(1 for row in s["rows"] for v in row if isinstance(v, dict))
    print(f"{ws.title}: rows={len(s['rows'])} cols={len(s['columns'])} formulas={n_formulas} "
          f"overrides={n_over} validations={len(s.get('validations', []))} "
          f"conditionals={len(s.get('conditionals', []))} freeze={s.get('freeze')} filter={s.get('autoFilter')}")

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(spec, f, ensure_ascii=False, indent=1)
print("json bytes:", len(open(OUT, "rb").read()))
